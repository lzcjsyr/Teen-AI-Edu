import openpyxl

XLSX = "/Users/dielangli/Desktop/AI培训/职场AI培训/8月12日培训/04_Wind基金筛选/初始池_基金筛选.xlsx"

# 1) 读取上一步的基金清单（来自「01_初始池」sheet：代码、名称、年化、规模）
wb = openpyxl.load_workbook(XLSX)
ws = wb["01_初始池"]
funds = []
for r in ws.iter_rows(min_row=2, values_only=True):
    code, name, ret, size = r[0], r[1], r[2], r[3]
    funds.append([code, name, ret, size])

# 2) 剔除关键词列表（命中即判定为被动/商品型）
KEYWORDS = ["ETF", "联接", "指数", "中证", "标普", "纳斯达克",
            "黄金", "上海金", "白银", "可转债", "期货"]

# 3) 判断规则：名称含任一关键词 → 剔除；其余保留（主动管理型）
def is_passive(name):
    return any(k in name for k in KEYWORDS)

retained = [f for f in funds if not is_passive(f[1])]
removed = [f for f in funds if is_passive(f[1])]

# 4) 输出保留清单
print(f"上一步清单总数 : {len(funds)}")
print(f"剔除（被动/商品型）: {len(removed)}")
print(f"保留（主动管理型）: {len(retained)}")
print("=" * 78)
print(f"{'基金代码':<14}{'基金名称':<40}{'近5年年化(%)':>14}{'规模(亿元)':>14}")
print("-" * 78)
for code, name, ret, size in retained:
    print(f"{str(code):<14}{str(name):<40}{ret:>14}{size:>14}")

# 5) 追加到 Excel，新增 sheet「02_主动管理型」
if "02_主动管理型" in wb.sheetnames:
    del wb["02_主动管理型"]
ws2 = wb.create_sheet("02_主动管理型")
ws2.append(["基金代码", "基金名称", "近5年年化收益率(%)", "规模(亿元)"])
for code, name, ret, size in retained:
    ws2.append([code, name, ret, size])
for i, w in enumerate([14, 40, 20, 14], start=1):
    ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
wb.save(XLSX)
print("=" * 78)
print(f"已追加 sheet「02_主动管理型」，共 {len(retained)} 只，文件已保存。")
