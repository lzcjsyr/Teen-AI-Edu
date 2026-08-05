#!/usr/bin/env python3
"""把一组图片按顺序合成为 16:9 PPTX。"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp"}


def natural_key(path: Path) -> list[object]:
    parts = re.split(r"(\d+)", path.stem)
    return [int(part) if part.isdigit() else part.lower() for part in parts]


def collect_images(image_dir: Path) -> list[Path]:
    images = [
        path
        for path in image_dir.iterdir()
        if path.is_file() and path.suffix.lower() in IMAGE_EXTS
    ]
    return sorted(images, key=natural_key)


def build_ppt(images: list[Path], out_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]
    for image_path in images:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            str(image_path),
            0,
            0,
            width=prs.slide_width,
            height=prs.slide_height,
        )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="把图片合成为图片版 PPTX")
    parser.add_argument("image_dir", type=Path, help="图片目录")
    parser.add_argument("--out", type=Path, required=True, help="输出 PPTX 文件")
    args = parser.parse_args()

    images = collect_images(args.image_dir)
    if not images:
        raise SystemExit(f"没有在目录中找到图片：{args.image_dir}")

    build_ppt(images, args.out)
    print(f"已生成 PPT：{args.out}，共 {len(images)} 页")


if __name__ == "__main__":
    main()

