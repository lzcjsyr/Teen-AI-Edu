#!/usr/bin/env python3
"""Excel 销售明细分类 — AI 驱动的 I/O 工具（灵活版）。

设计原则：脚本不做任何"理解/判断"，只负责读、写、提取。
所有格式理解（表头在哪一行、品名在哪一列、新列插在哪）都由 AI 完成，
AI 把结论通过参数/映射 JSON 传给脚本，脚本忠实地执行。

子命令：
  read   — 把整张表的"形状"和前若干行原样dump给 AI，供 AI 理解格式
  names  — 按 AI 指定的 表头行 + 品名列，提取唯品名（全量扫描）
  write  — 按 AI 给的映射 JSON，把分类结果写回 Excel
"""

import io, json, sys, re, argparse
from pathlib import Path
from dataclasses import dataclass
from typing import Protocol, cast

import importlib

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
import openpyxl


# xlrd 无类型桩，这里用最小 Protocol 描述脚本实际用到的接口，避免 import_module 返回 Any 的告警。
class _XlrdCell(Protocol):
    value: object

class _XlrdSheet(Protocol):
    nrows: int
    ncols: int
    def cell(self, row: int, col: int) -> _XlrdCell: ...

class _XlrdBook(Protocol):
    def sheet_by_index(self, n: int) -> _XlrdSheet: ...

class _XlrdModule(Protocol):
    def open_workbook(self, filename: str, logfile: object = ...) -> _XlrdBook: ...


# ---------------------------------------------------------------------------
# 品名标准化 — 去掉规格/单位/括号，方便 AI 分类时聚焦品类
# ---------------------------------------------------------------------------

def normalize_name(name: str) -> str:
    text = name.strip()
    text = re.sub(r"\s+", "", text)
    text = re.sub(r"^[（(][^）)]+[）)]", "", text)
    prev = ""
    while text != prev:
        prev = text
        text = re.sub(
            r"^(.*)[xX*×/\-\d\s]+(?:kg|g|ml|l|升|只|个|支|瓶|桶|包|件|袋|盒|箱|头|两|米|张|盒)$",
            r"\1", text, flags=re.IGNORECASE,
        )
        text = re.sub(
            r"^(.*)\d+(?:kg|g|ml|l|升|只|个|支|瓶|桶|包|件|袋|盒|箱|头|两|米|张|盒)$",
            r"\1", text, flags=re.IGNORECASE,
        )
        text = re.sub(r"^(.*)/(?:盒|包|件|瓶|桶|扎|双|把|斤|袋|箱|个|只|支|盒)$", r"\1", text)
        text = re.sub(r"^(.*)[xX*×/]\d+$", r"\1", text)
    text = re.sub(
        r"\d+(?:\.\d+)?\s*(?:kg|g|ml|l|升|只|个|支|瓶|桶|包|件|袋|盒|箱|头|两|米|张|盒)",
        "", text, flags=re.IGNORECASE,
    )
    text = re.sub(r'[（()）\-+*#/"]', "", text)
    return text


# ---------------------------------------------------------------------------
# Excel 读取（抑制 xlrd 噪音）
# ---------------------------------------------------------------------------

def _read(path: Path) -> list[list[object]]:
    ext = path.suffix.lower()
    if ext == ".xls":
        try:
            xlrd_mod = cast(_XlrdModule, cast("object", importlib.import_module("xlrd")))
        except ImportError:
            raise RuntimeError("读取 .xls 需要 xlrd，请先安装：pip install xlrd") from None
        # xlrd 的 open_workbook 在函数定义时就把 logfile 默认绑定成了当时的 sys.stdout，
        # 因此仅重定向 sys.stdout 无法拦截 "OLE2 inconsistency" 一类告警，
        # 必须显式传入 logfile，否则告警会混进 stdout 污染 JSON 输出。
        sink = io.StringIO()
        old_out, old_err = sys.stdout, sys.stderr
        sys.stdout = sys.stderr = sink
        try:
            wb = xlrd_mod.open_workbook(str(path), logfile=sink)
        finally:
            sys.stdout, sys.stderr = old_out, old_err
        sheet = wb.sheet_by_index(0)
        return [[sheet.cell(r, c).value for c in range(sheet.ncols)]
                for r in range(sheet.nrows)]
    elif ext == ".xlsx":
        wb = openpyxl.load_workbook(path, data_only=True)
        sheet = wb.active
        return [list(row) for row in sheet.iter_rows(values_only=True)] if sheet else []
    else:
        raise ValueError(f"不支持的文件类型：{ext}（仅支持 .xls / .xlsx）")


def _cell(v: object) -> object:
    """把单元格值转成对 AI 友好的 JSON 形式。"""
    if v is None:
        return None
    if isinstance(v, float):
        # 整数形态的小数（如 12.0）转成 int，避免 AI 困惑
        return int(v) if v.is_integer() else v
    if isinstance(v, (int,)):
        return v
    return str(v).strip()


# ---------------------------------------------------------------------------
# read — 把表"形状"和样本交给 AI，由 AI 判断格式（脚本不猜列）
# ---------------------------------------------------------------------------

PREVIEW_LIMIT = 60  # 给 AI 看的样本行数上限


def cmd_read(input_file: str) -> int:
    path = Path(input_file).expanduser().resolve()
    if not path.is_file():
        print(f"ERROR: 文件不存在 {path}", file=sys.stderr)
        return 1

    rows = _read(path)
    total_cols = max((len(r) for r in rows), default=0)

    # 仅作为"提示"输出，AI 可采纳也可无视；脚本不据此做任何硬性判断
    suggestions = _suggest_header(rows)

    preview = [
        {"row_index": i, "cells": [_cell(c) for c in rows[i]]}
        for i in range(min(len(rows), PREVIEW_LIMIT))
    ]

    result = {
        "input_file": str(path),
        "total_rows": len(rows),
        "total_cols": total_cols,
        "suggestions": suggestions,
        "preview_rows": preview,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


# 品名列关键词，按特异性从高到低分层。
# 分层的原因：像“单位名称”“客户名称”这类列也含“名称”，若只做宽松的子串匹配，
# 会先命中它们而错过真正的“商品名称”。因此先扫一遍强信号，再退回弱信号。
_NAME_KEYWORDS: tuple[tuple[str, ...], ...] = (
    ("商品名称", "货物名称", "物料名称", "存货名称", "产品名称", "品名"),
    ("商品", "货物", "物料", "存货"),
    ("名称",),
)


def _suggest_header(rows: list[list[object]]) -> dict[str, int]:
    """尽力给 AI 一个起点建议；找不到就返回 -1，绝不报错。

    注意：这里只是"提示"，最终的表头行与品名列必须由 AI 依据样本判断。
    弱关键词很容易误命中"单位名称"一类列，不可直接采信。
    """
    for i, row in enumerate(rows):
        if not row:
            continue
        if len([c for c in row if c is not None and str(c).strip()]) < 2:
            continue
        for tier in _NAME_KEYWORDS:
            hit = next((j for j, c in enumerate(row)
                        if c is not None and any(k in str(c) for k in tier)), -1)
            if hit >= 0:
                return {"header_row": i, "name_col": hit}
    return {"header_row": -1, "name_col": -1}


# ---------------------------------------------------------------------------
# names — 按 AI 指定的 表头行 + 品名列，提取唯品名（全量扫描）
# ---------------------------------------------------------------------------

def cmd_names(input_file: str, header_idx: int, name_col: int) -> int:
    path = Path(input_file).expanduser().resolve()
    if not path.is_file():
        print(f"ERROR: 文件不存在 {path}", file=sys.stderr)
        return 1
    if header_idx < 0 or name_col < 0:
        print("ERROR: 必须提供 --header-idx 与 --name-col", file=sys.stderr)
        return 1

    rows = _read(path)
    if header_idx >= len(rows):
        print(f"ERROR: header_idx={header_idx} 超出总行数 {len(rows)}", file=sys.stderr)
        return 1

    seen: dict[str, str] = {}  # normalized → original
    for r in range(header_idx + 1, len(rows)):
        row = rows[r]
        if not row or name_col >= len(row):
            continue
        name = str(row[name_col]).strip() if row[name_col] is not None else ""
        if not name:
            continue
        n = normalize_name(name)
        if n and n not in seen:
            seen[n] = name

    result = {
        "input_file": str(path),
        "header_idx": header_idx,
        "col_name_idx": name_col,
        "data_rows": len(rows) - header_idx - 1,
        "unique_count": len(seen),
        "unique_names": [
            {"normalized": n, "original": o} for n, o in seen.items()
        ],
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


# ---------------------------------------------------------------------------
# write — 按 AI 给的映射 JSON 回写分类结果
# ---------------------------------------------------------------------------

def cmd_write(input_file: str, output_file: str, map_file: str) -> int:
    input_path = Path(input_file).expanduser().resolve()
    output_path = Path(output_file).expanduser().resolve()
    map_path = Path(map_file).expanduser().resolve()

    if not input_path.is_file():
        print(f"ERROR: 输入文件不存在 {input_path}", file=sys.stderr)
        return 1
    if not map_path.is_file():
        print(f"ERROR: 分类映射文件不存在 {map_path}", file=sys.stderr)
        return 1

    with open(map_path, encoding="utf-8") as f:
        mapping = json.load(f)

    classifications: dict[str, dict[str, str]] = mapping.get("classifications", {})
    header_idx: int = mapping.get("header_idx", -1)
    col_name: int = mapping.get("col_name_idx", -1)

    if not classifications or header_idx < 0 or col_name < 0:
        print("ERROR: 分类映射格式不正确，缺少 classifications / header_idx / col_name_idx",
              file=sys.stderr)
        return 1

    rows = _read(input_path)

    # 新列插入位置：优先用 AI 显式指定的 insert_idx；否则退化为"备注列之前/行尾"
    if "insert_idx" in mapping and mapping["insert_idx"] is not None:
        insert_idx = int(mapping["insert_idx"])
    else:
        col_remark = mapping.get("col_remark_idx", -1)
        insert_idx = col_remark if col_remark != -1 else len(rows[header_idx])

    out_rows: list[tuple[list[object], bool]] = []
    header_font = Font(bold=True)

    # 表头之前的行
    for r in range(header_idx):
        old = list(rows[r])
        out_rows.append((old[:insert_idx] + ["", ""] + old[insert_idx:], False))

    # 表头行
    old = list(rows[header_idx])
    out_rows.append((old[:insert_idx] + ["科目代码", "科目名称"] + old[insert_idx:], True))

    # 数据行
    for r in range(header_idx + 1, len(rows)):
        old = list(rows[r])
        if not old:
            continue
        non_empty = [c for c in old if c is not None and str(c).strip() != ""]
        if not non_empty:
            out_rows.append((old[:insert_idx] + ["", ""] + old[insert_idx:], False))
            continue

        name = str(old[col_name]).strip() if col_name < len(old) else ""
        if name:
            key = normalize_name(name)
            cls = classifications.get(key, {})
            code, subject = cls.get("code", ""), cls.get("subject", "")
        else:
            code, subject = "", ""

        out_rows.append((old[:insert_idx] + [code, subject] + old[insert_idx:], False))

    # 写入
    wb = Workbook()
    ws = wb.active
    assert ws is not None, "新建的 Workbook 必然至少含一个工作表"
    ws.title = "已分类明细"

    for r_idx, (row_data, is_header) in enumerate(out_rows, start=1):
        ws.append(row_data)
        if is_header:
            for cell in ws[r_idx]:
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = f"A{header_idx + 2}"

    # 用列序号 + get_column_letter 生成列字母，避免依赖 col_cells[0]（可能为 MergedCell）
    for col_idx, col_cells in enumerate(ws.columns, start=1):
        max_len = max((len(str(c.value or "")) for c in col_cells), default=0)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 28)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"XLSX={output_path}")
    return 0


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(description="Excel 分类 I/O 工具（AI 驱动·灵活版）")
    sub = parser.add_subparsers(dest="cmd", required=True)

    p_read = sub.add_parser("read", help="dump 表结构给 AI 理解（不做任何判断）")
    _ = p_read.add_argument("input_file")

    p_names = sub.add_parser("names", help="按 AI 指定的表头行+品名列提取唯品名")
    _ = p_names.add_argument("input_file")
    _ = p_names.add_argument("--header-idx", type=int, required=True)
    _ = p_names.add_argument("--name-col", type=int, required=True)

    p_write = sub.add_parser("write", help="按 AI 映射 JSON 回写分类结果")
    _ = p_write.add_argument("input_file")
    _ = p_write.add_argument("--output", required=True)
    _ = p_write.add_argument("--map", required=True)

    # 用结构化命名空间接收解析结果：_CliArgs 是 Namespace 的子类且字段带默认值，
    # 因此 args.* 具有确定类型（替代 argparse 默认的 Any 属性），且无需 cast。
    @dataclass
    class _CliArgs(argparse.Namespace):
        cmd: str = ""
        input_file: str = ""
        header_idx: int = 0
        name_col: int = 0
        output: str = ""
        map: str = ""

    ns = _CliArgs()
    _ = parser.parse_args(namespace=ns)
    args = ns

    if args.cmd == "read":
        return cmd_read(args.input_file)
    elif args.cmd == "names":
        return cmd_names(args.input_file, args.header_idx, args.name_col)
    elif args.cmd == "write":
        return cmd_write(args.input_file, args.output, args.map)
    return 1


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as e:
        print(f"ERROR: {e}", file=sys.stderr)
        sys.exit(1)
